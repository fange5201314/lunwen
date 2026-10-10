"""Project source-level priority audit into reviewable inventory and adapters."""
from pathlib import Path
import csv
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
audit = json.loads((ROOT / 'research/111-resource-audit.json').read_text())
evidence = json.loads((ROOT / 'research/111-repo-evidence.json').read_text())
resources = audit['resources']
keys = ['id', 'name', 'source_index_id', 'recipes', 'repos', 'status', 'use', 'input', 'output', 'validation', 'gaps']
def flat(value):
    if isinstance(value, list):
        return '；'.join(str(x) for x in value)
    return '' if value is None else value
rows = [{k: flat(item.get(k)) for k in keys} for item in resources]
with (ROOT / '111重点资源融合清单.csv').open('w', newline='', encoding='utf-8-sig') as handle:
    writer = csv.DictWriter(handle, fieldnames=keys, lineterminator='\n')
    writer.writeheader()
    writer.writerows(rows)

adapters = []
for repo in evidence['repositories']:
    source_matches = [r for r in resources if repo['repo'] in r['repos']]
    adapters.append({
        'repo': repo['repo'], 'url': 'https://github.com/' + repo['repo'],
        'source_ids': [r['id'] for r in source_matches],
        'module': '；'.join(dict.fromkeys(r['use'] for r in source_matches)),
        'recipes': list(dict.fromkeys(s for r in source_matches for s in r['recipes'])),
        'commit': repo.get('commit'), 'license': repo.get('license'),
        'readme_url': repo.get('readme_url'), 'runtime_verified': False,
        'reuse': '固定版本协议/方法或候选接口；具体执行接缝待自测',
        'validation': '；'.join(dict.fromkeys(r['validation'] for r in source_matches)),
        'contracts': [c for c in evidence['reusable_contracts'] if c['repo'] == repo['repo']]
    })
manifest_path = ROOT / 'config/integration-manifest.json'
manifest = json.loads(manifest_path.read_text())
manifest['design_version'] = '2.1'
manifest['priority_111_modules'] = adapters
manifest['fusion_layer'] = {
    'blueprints': 'fusion-blueprints.json', 'gates': 'fusion-gate-policy.json',
    'evidence_contract': 'evidence-impact.schema.json',
    'runtime_verified': False, 'compile_owner': 'thesis-domain-service',
    'no_additional_business_state_owner': True
}
for item in manifest['modules']:
    direct = [r['id'] for r in resources if any(m['repo'].lower() == item['repo'].lower() and
              m['relation'] == 'direct_source_repo' for m in r['github_matches'])]
    item['priority_111_source_ids'] = direct
manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
adapter_keys = ['repo', 'source_ids', 'module', 'recipes', 'commit', 'license', 'runtime_verified', 'reuse', 'validation']
with (ROOT / '111补充开源适配表.csv').open('w', newline='', encoding='utf-8-sig') as handle:
    writer = csv.DictWriter(handle, fieldnames=adapter_keys, lineterminator='\n')
    writer.writeheader()
    writer.writerows({k: flat(a.get(k)) if k != 'runtime_verified' else False for k in adapter_keys} for a in adapters)

book_path = ROOT / '资源复用与缺口清单.xlsx'
book = load_workbook(book_path)
def write_sheet(name, columns, records):
    if name in book:
        del book[name]
    sheet = book.create_sheet(name)
    sheet.append(columns)
    for row in records:
        sheet.append([flat(row.get(k)) for k in columns])
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='665395')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 30 if col[0].value not in ['id', 'source_index_id'] else 14
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
write_sheet('111重点融合', keys, rows)
write_sheet('111开源适配', adapter_keys, adapters)
contract_keys = ['repo', 'commit', 'path', 'line_start', 'line_end', 'section', 'short_summary', 'local_mapping', 'file_url']
write_sheet('协议复用证据', contract_keys, evidence['reusable_contracts'])
book.save(book_path)

doc = ['# 111 资源协议复用证据', '', '以下为固定提交核心文件的只读核对，方法映射属于工作站设计；执行接入尚未验证。', '']
for i, c in enumerate(evidence['reusable_contracts'], 1):
    doc += [f"## {i}. {c['repo']} · {c['section']}", '',
            f"来源：[{c['path']}]({c['file_url']})，第 {c['line_start']}–{c['line_end']} 行；提交 `{c['commit']}`。", '',
            c['short_summary'], '', '工作站落点：' + c['local_mapping'], '']
(ROOT / '111资源协议复用证据.md').write_text('\n'.join(doc))
assert len(resources) == 23
assert len(adapters) == 14
print('23 priority records, 14 supplementary repositories, 12 protocol contracts projected.')
