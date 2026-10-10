"""Repair escaped Markdown pipes and refresh all public catalog projections.

The original row is preserved for provenance. Run after source audits are updated.
"""
from pathlib import Path
import csv
import json
import re
from urllib.parse import unquote
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
path = ROOT / 'config/resource-catalog.json'
catalog = json.loads(path.read_text())
changes = []
for item in catalog:
    parts = re.split(r'(?<!\\)\|', item.get('raw_row', '').strip().strip('|'))
    if item['id'].startswith('CL-') and len(parts) == 5:
        for key, value in zip(('name', 'summary', 'interfaces', 'evidence'), parts[1:]):
            value = value.strip().replace('\\|', '|')
            if value != item.get(key):
                changes.append({'id': item['id'], 'field': key})
                item[key] = value
    item['priority_111'] = unquote(item.get('url', '')).split('/')[-1].startswith('111')
path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2))

js_path = ROOT / 'config/catalog-data.js'
raw = js_path.read_text()
data = json.loads(raw.removeprefix('window.CATALOG=').rstrip().rstrip(';'))
data['resources'] = [{k: v for k, v in row.items() if k != 'raw_row'} for row in catalog]
encoded = json.dumps(data, ensure_ascii=False).replace('</', '<\\/')
js_path.write_text('window.CATALOG=' + encoded + ';\n')
prototype = ROOT / '可点击原型.html'
text = prototype.read_text()
start = text.index('<script>const CATALOG=') + len('<script>const CATALOG=')
end = text.index('; const CURATED=', start)
prototype.write_text(text[:start] + encoded + text[end:])

keys = ['id', 'name', 'category', 'summary', 'interfaces', 'evidence', 'url']
with (ROOT / '资源全量索引.csv').open('w', newline='', encoding='utf-8-sig') as handle:
    writer = csv.DictWriter(handle, fieldnames=keys, extrasaction='ignore', lineterminator='\n')
    writer.writeheader()
    writer.writerows(catalog)
workbook_path = ROOT / '资源复用与缺口清单.xlsx'
workbook = load_workbook(workbook_path)
sheet = workbook['全量资源']
sheet.delete_rows(2, sheet.max_row)
for item in catalog:
    sheet.append([item.get(key, '') for key in keys])
sheet.freeze_panes = 'A2'
sheet.auto_filter.ref = sheet.dimensions
workbook.save(workbook_path)
report = {'records': len(catalog), 'priority_111_records': sum(r['priority_111'] for r in catalog),
          'corrections': changes, 'rule': 'split on unescaped Markdown pipe; preserve original raw_row'}
(ROOT / 'research/catalog-parser-repair.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
assert len(catalog) == 884
assert report['priority_111_records'] == 23
by_id = {r['id']: r for r in catalog}
assert 'Rovai 0.3' in by_id['CL-0015']['name']
assert by_id['CL-0006']['interfaces'] == 'API'
print(json.dumps({'records': 884, 'priority_records': 23, 'corrected_fields': len(changes)}))
