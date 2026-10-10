from pathlib import Path
import html,re,json
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).parent.parent

def inline(s):
 s=html.escape(s)
 s=re.sub(r'`([^`]+)`',r'<code>\1</code>',s)
 s=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',s)
 s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'<a href="\2">\1</a>',s)
 return s

def convert(text):
 lines=text.splitlines();out=[];i=0;toc=[]
 while i<len(lines):
  line=lines[i]
  if not line.strip():i+=1;continue
  if line.startswith('```'):
   code=[];i+=1
   while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
   out.append('<pre><code>'+html.escape('\n'.join(code))+'</code></pre>');i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    row=[s.strip() for s in lines[i].strip().strip('|').split('|')]
    if not all(re.match(r'^:?-+:?$',x) for x in row):rows.append(row)
    i+=1
   out.append('<div class="table"><table><thead><tr>'+''.join('<th>'+inline(s)+'</th>' for s in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline(s)+'</td>' for s in row)+'</tr>' for row in rows[1:])+'</tbody></table></div>');continue
  m=re.match(r'^(#{1,6}) (.+)',line)
  if m:
   level=len(m[1]);label=m[2];slug='s'+str(len(out));
   if level==2:toc.append((slug,label))
   out.append(f'<h{level} id="{slug}">{inline(label)}</h{level}>');i+=1;continue
  if re.match(r'^(- |\d+\. )',line):
   numbered=bool(re.match(r'^\d+\. ',line));kind='ol' if numbered else 'ul';items=[]
   while i<len(lines) and re.match(r'^(- |\d+\. )',lines[i]):items.append(re.sub(r'^(- |\d+\. )','',lines[i]));i+=1
   out.append('<'+kind+'>'+''.join('<li>'+inline(s)+'</li>' for s in items)+'</'+kind+'>');continue
  para=[line];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- |\d+\. )',lines[i]):para.append(lines[i]);i+=1
  out.append('<p>'+inline(' '.join(para))+'</p>')
 return '\n'.join(out),toc
style='''body{margin:0;color:#302d43;background:#f5f4f9;font:14px/1.9 "Noto Sans CJK SC","Microsoft YaHei",sans-serif}main{max-width:1000px;margin:45px auto;background:white;padding:60px 65px;border:1px solid #e8e4f0;border-radius:17px}.tag{color:#8a76b3;font-size:11px;letter-spacing:2px}.toplink{font-size:12px;text-align:right}h1{font-size:30px;line-height:1.5;margin:12px 0 25px}h2{font-size:22px;line-height:1.5;padding-top:25px;margin-top:36px;border-top:1px solid #eae5f2;color:#675087}h3{font-size:16px;margin-top:28px}p{margin:12px 0}table{border-collapse:collapse;width:100%;font-size:11px;line-height:1.7;table-layout:fixed}th{background:#ebe7f5;color:#655280;font-weight:600}td,th{text-align:left;vertical-align:top;padding:11px;border:1px solid #e7e2ef;word-break:break-word}td p{margin:0}tr:nth-child(even){background:#fbfafc}.table{margin:20px 0}pre{font-size:11px;line-height:1.7;background:#f2eff7;border:1px solid #e5dfef;border-radius:9px;padding:18px;overflow:auto;white-space:pre-wrap;word-break:break-word}code{font-family:ui-monospace,monospace}a{color:#7561a5;text-decoration:none}.toc{padding:20px 25px;background:#f7f5fb;border-radius:12px;font-size:12px}.toc a{display:block}ul,ol{padding-left:23px}li{margin:6px 0}.covernote{padding:15px 20px;border-left:3px solid #9e8ac9;background:#f5f2fa;font-size:12px;color:#776988}figure{margin:30px 0;text-align:center}figure img{max-width:100%;height:auto}figcaption{font-size:11px;color:#8b829c;margin-top:12px}@media(max-width:800px){main{padding:25px;margin:15px;border-radius:10px}h1{font-size:25px}}@media print{@page{size:A4;margin:18mm 16mm}body{background:white;font-size:10pt;line-height:1.8}main{border:0;border-radius:0;padding:0;margin:0;max-width:none}.toplink,.toc{display:none}h1{font-size:24pt}h2{font-size:16pt;break-after:avoid}h3{font-size:12pt;break-after:avoid}table{font-size:8pt}tr{break-inside:avoid}thead{display:table-header-group}pre{font-size:8pt}figure{break-inside:avoid}figure img{max-height:235mm;object-fit:contain}.print-cover{break-after:page}}'''
for name in ['架构设计与方案说明','资源清单与使用说明']:
 text=(root/(name+'.md')).read_text();body,toc=convert(text)
 figures=''
 if name=='架构设计与方案说明':
  figures='''<h2>附图：系统与协作设计</h2><figure><img src="架构图预览.png"><figcaption>系统架构：Archify实际生成；设计拓扑，不代表运行事实。</figcaption></figure><figure><img src="多Agent协作.svg"><figcaption>多Agent岗位协作：搜索、方案、讨论、确认、交接、独立检查与用户验收。</figcaption></figure><figure><img src="失败与恢复.svg"><figcaption>请求级故障转移与人工任务恢复的边界。</figcaption></figure>'''
 doc='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+name+'</title><style>'+style+'</style><main><div class="toplink"><a href="可点击原型.html">打开交互原型 ↗</a>　<a href="架构图.html">探索架构图 ↗</a></div><div class="tag">THESIS ATELIER / DESIGN REVIEW v2.0</div><div class="covernote">真实交付：资源、设计、交互原型与图。示例执行与论文产物为模拟，专业工具待接入验证。</div><div class="toc"><strong>阅读目录</strong>'+''.join('<a href="#'+slug+'">'+html.escape(label)+'</a>' for slug,label in toc)+'</div>'+body+figures+'</main></html>'
 (root/(name+'.html')).write_text(doc)
# Editable Word version of the complete proposal
text=(root/'架构设计与方案说明.md').read_text();doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
 styleobj=doc.styles[name];styleobj.font.name='Microsoft YaHei';styleobj._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');styleobj.font.color.rgb=RGBColor.from_string('675087')
lines=text.splitlines();i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('```'):
  code=[];i+=1
  while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
  p=doc.add_paragraph('\n'.join(code));p.style='Normal'
  for run in p.runs:run.font.name='Consolas';run.font.size=Pt(8)
  i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   row=[s.strip() for s in lines[i].strip().strip('|').split('|')]
   if not all(re.match(r'^:?-+:?$',x) for x in row):rows.append(row)
   i+=1
  table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
  for j,s in enumerate(rows[0]):table.rows[0].cells[j].text=s
  # Repeat header on every page.
  trPr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trPr.append(repeat)
  for row in rows[1:]:
   cells=table.add_row().cells
   for j,s in enumerate(row):cells[j].text=s
  for row in table.rows:
   for cell in row.cells:
    for p in cell.paragraphs:
     for run in p.runs:run.font.size=Pt(8)
  doc.add_paragraph();continue
 m=re.match(r'^(#{1,6}) (.+)',line)
 if m:doc.add_heading(m[2],level=0 if len(m[1])==1 else min(3,len(m[1])-1));i+=1;continue
 if line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet')
 elif re.match(r'^\d+\. ',line):doc.add_paragraph(re.sub(r'^\d+\. ','',line),style='List Number')
 else:doc.add_paragraph(line)
 i+=1
footer=sec.footer.paragraphs[0];footer.alignment=2;run=footer.add_run('THESIS ATELIER · 设计评审 v2.0　');run.font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(root/'架构设计与方案说明.docx')
print('Proposal HTML, resource HTML and DOCX generated')
